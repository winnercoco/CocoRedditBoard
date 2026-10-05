from pathlib import Path
from openpyxl import load_workbook, Workbook


BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"

MASTER_FILE = DATA_DIR / "2-master_reddit_links_curated.xlsx"
OUTPUT_FILE = DATA_DIR / "3-reddit_links.xlsx"


def main():

    print("=" * 60)
    print("Creating links.xlsx")
    print("=" * 60)

    # ---------------------------------------------------------
    # Load master workbook using calculated values.
    #
    # IMPORTANT:
    # data_only=True means formula cells return their
    # cached/calculated result instead of the formula itself.
    # ---------------------------------------------------------

    wb = load_workbook(
        MASTER_FILE,
        data_only=True
    )

    ws = wb.active

    print(f"Source sheet : {ws.title}")

    # ---------------------------------------------------------
    # Find Flag column
    # ---------------------------------------------------------

    flag_column = None

    for cell in ws[1]:

        if (
            cell.value is not None
            and str(cell.value).strip().lower() == "flag"
        ):
            flag_column = cell.column
            break

    if flag_column is None:
        raise ValueError("Could not find the Flag column.")

    print(f"Flag column  : {flag_column}")

    # ---------------------------------------------------------
    # Create completely clean workbook
    # ---------------------------------------------------------

    output_wb = Workbook()
    output_ws = output_wb.active
    output_ws.title = ws.title

    # ---------------------------------------------------------
    # Copy header as VALUES only
    # ---------------------------------------------------------

    for column in range(1, ws.max_column + 1):

        output_ws.cell(
            row=1,
            column=column
        ).value = ws.cell(
            row=1,
            column=column
        ).value

    # ---------------------------------------------------------
    # Copy rows whose CALCULATED Flag value is "done"
    # ---------------------------------------------------------

    output_row = 2
    copied_count = 0

    for source_row in range(2, ws.max_row + 1):

        flag_value = ws.cell(
            row=source_row,
            column=flag_column
        ).value

        if (
            flag_value is None
            or str(flag_value).strip().lower() != "done"
        ):
            continue

        # Copy VALUES ONLY.
        #
        # No formulas.
        # No formatting.
        # No hyperlinks.
        # No comments.
        # No workbook structures.
        #
        for column in range(1, ws.max_column + 1):

            output_ws.cell(
                row=output_row,
                column=column
            ).value = ws.cell(
                row=source_row,
                column=column
            ).value

        output_row += 1
        copied_count += 1

    # ---------------------------------------------------------
    # Save
    # ---------------------------------------------------------

    output_wb.save(OUTPUT_FILE)

    wb.close()
    output_wb.close()

    # ---------------------------------------------------------
    # Report
    # ---------------------------------------------------------

    print()
    print("=" * 60)
    print("Finished")
    print("=" * 60)
    print(f"Rows copied : {copied_count}")
    print(f"Total rows  : {copied_count + 1} including header")
    print(f"Output      : {OUTPUT_FILE}")
    print("=" * 60)


if __name__ == "__main__":
    main()